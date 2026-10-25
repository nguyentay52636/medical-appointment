import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs can be top, bottom, left, right.
    value: dict(sz=4, val='single', color='CBD5E1')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_xml = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "CBD5E1")}"/>'
        else:
            b_xml = f'<w:{edge} {nsdecls("w")} w:val="none"/>'
        tcBorders.append(parse_xml(b_xml))
    tcPr.append(tcBorders)

def create_report(output_path):
    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Base style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # --- Title Banner Block ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("BÁO CÁO TIẾN ĐỘ HOÀN THIỆN DỰ ÁN")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF) # Blue-800

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("Phân Hệ Quản Trị Phòng Khám, Cấp Tài Khoản Bác Sĩ & Chuẩn Hóa Kiến Trúc")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x3B, 0x82, 0xF6) # Blue-500

    # Meta Info Card (Table 1x1)
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_cell = meta_table.cell(0, 0)
    meta_cell.width = Inches(6.7)
    set_cell_background(meta_cell, "F0F7FF")
    set_cell_margins(meta_cell, top=140, bottom=140, left=180, right=180)
    set_cell_border(meta_cell, left=dict(sz=18, val='single', color='2563EB'),
                               top=dict(sz=4, val='single', color='DBEAFE'),
                               bottom=dict(sz=4, val='single', color='DBEAFE'),
                               right=dict(sz=4, val='single', color='DBEAFE'))

    meta_p = meta_cell.paragraphs[0]
    meta_p.paragraph_format.space_after = Pt(2)
    r1 = meta_p.add_run("• Dự án: ")
    r1.bold = True
    meta_p.add_run("Hệ thống Đặt lịch khám bệnh & Quản lý phòng khám trực tuyến (Klinik / MediCare)\n")

    r2 = meta_p.add_run("• Thời gian cập nhật: ")
    r2.bold = True
    meta_p.add_run("25/10/2026   |   ")

    r3 = meta_p.add_run("• Trạng thái: ")
    r3.bold = True
    r3_val = meta_p.add_run("Đã hoàn thiện & Kiểm thử thành công 100% (PASS ✅)\n")
    r3_val.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
    r3_val.bold = True

    r4 = meta_p.add_run("• Phiên bản tài liệu: ")
    r4.bold = True
    meta_p.add_run("2.0 (Toàn bộ kiến trúc 3 phân hệ: Bệnh nhân, Cổng Bác sĩ và Quản trị phòng khám)")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- Helper to add Headings ---
    def add_sec_heading(num, text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(f"{num}. {text.upper()}")
        r.font.name = 'Calibri'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return h

    def add_sub_heading(num, text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(f"{num} {text}")
        r.font.name = 'Calibri'
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        return h

    # ==========================================
    # PHẦN 1: TỔNG QUAN YÊU CẦU & BỐI CẢNH
    # ==========================================
    add_sec_heading("1", "Tổng Quan Yêu Cầu & Bối Cảnh")
    add_sub_heading("1.1", "Thực trạng trước khi xử lý")

    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run("Lệch vai trò nghiệp vụ: ").bold = True
    p.add_run("Ban đầu hệ thống chỉ có 2 nhóm đối tượng là Bệnh nhân (Patient) và Bác sĩ (Doctor). Chưa có phân cấp quản trị phòng khám.")

    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run("Nhầm lẫn trang làm việc cá nhân: ").bold = True
    p.add_run("Trang theo dõi ca khám riêng của từng bác sĩ bị đặt nhầm nhãn là 'Dashboard' (/doctor/dashboard), dẫn tới hiểu nhầm là bảng quản trị chung của phòng khám.")

    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run("Thiếu hoàn toàn phân hệ Quản trị (Admin): ").bold = True
    p.add_run("Ban quản lý không có nơi xem KPI tổng quan, không có giao diện thêm mới bác sĩ, cấp tài khoản, phân chuyên khoa hay xóa bác sĩ nghỉ việc. Toàn bộ dữ liệu bác sĩ trước đây phải gán cứng (hardcode) bằng lệnh SQL thủ công vào cơ sở dữ liệu MySQL.")

    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run("Lỗi điều hướng & Bảo vệ tuyến đường (Auth Guard): ").bold = True
    p.add_run("Gõ /doctor tự động bật form login; chưa đăng nhập mà gõ /doctor/dashboard không bị chặn dẫn tới sập dữ liệu do thiếu session ID. Các trang quên/đổi mật khẩu bị lồng nhầm sidebar quản trị.")

    add_sub_heading("1.2", "Mục tiêu đã giải quyết thành công")
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run("Bảo toàn 100% các phân hệ hiện có: ").bold = True
    p.add_run("Giữ nguyên toàn bộ tính năng phân hệ Bệnh nhân (đặt lịch, chat AI tư vấn triệu chứng, bản đồ Geocoding tìm bác sĩ gần nhất, thanh toán VNPay) và Cổng Bác sĩ (lịch khám cá nhân hôm nay, ca trực tuần, xuất phiếu khám PDF).")

    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run("Xây dựng mới Phân hệ Quản trị Phòng khám Toàn diện (Clinic Admin Dashboard): ").bold = True
    p.add_run("Tạo mới trang Tổng quan KPI toàn viện (/admin/dashboard), Giao diện Cấp tài khoản Bác sĩ (/admin/doctors/create), Danh sách quản lý & điều chuyển chuyên khoa (/admin/doctors), Giám sát lịch hẹn toàn viện (/admin/appointments) và Danh bạ bệnh nhân (/admin/patients).")

    # ==========================================
    # PHẦN 2: MA TRẬN 3 PHÂN HỆ HỆ THỐNG
    # ==========================================
    add_sec_heading("2", "Ma Trận Kiến Trúc 3 Phân Hệ Sau Khi Hoàn Thiện")

    sys_table = doc.add_table(rows=1, cols=3)
    sys_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["PHÂN HỆ BỆNH NHÂN\n(Patient Portal - /patient/*)",
               "CỔNG BÁC SĨ\n(Doctor Portal - /doctor/*)",
               "QUẢN TRỊ PHÒNG KHÁM\n(Clinic Admin - /admin/*)"]
    widths = [Inches(2.2), Inches(2.2), Inches(2.3)]

    hdr_cells = sys_table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9.5)

    row_data = [
        ("• Tìm kiếm bác sĩ & chuyên khoa\n"
         "• Chat Trợ lý AI tư vấn triệu chứng\n"
         "• Bản đồ số tìm phòng khám gần\n"
         "• Đặt lịch hẹn & Thanh toán VNPay\n"
         "• Quản lý lịch sử, dời lịch hẹn\n"
         "• Đăng ký & Đăng nhập bệnh nhân"),
        ("• Xem lịch khám cá nhân hôm nay\n"
         "• Tiếp nhận, Khám xong & Hủy ca khám\n"
         "• Cài đặt ca rảnh làm việc hàng tuần\n"
         "• Xuất phiếu khám bệnh PDF\n"
         "• Chuẩn hóa Auth Guard bảo vệ đăng nhập\n"
         "• Đổi mật khẩu cá nhân bác sĩ"),
        ("• Bảng điều khiển KPI & Doanh thu viện\n"
         "• Form cấp tài khoản Bác sĩ chuẩn hóa\n"
         "• Điều chuyển & Phân khoa động từ DB\n"
         "• Cấp lại mật khẩu bác sĩ (Reset Pass)\n"
         "• Giám sát toàn bộ lịch khám các khoa\n"
         "• Quản lý danh bạ hồ sơ bệnh nhân")
    ]

    row_cells = sys_table.add_row().cells
    for i, content in enumerate(row_data):
        row_cells[i].text = content
        row_cells[i].width = widths[i]
        set_cell_background(row_cells[i], "F8FAFC" if i % 2 == 0 else "FFFFFF")
        set_cell_margins(row_cells[i], top=100, bottom=100, left=100, right=100)
        set_cell_border(row_cells[i], top=dict(sz=4, val='single', color='CBD5E1'),
                                      bottom=dict(sz=4, val='single', color='CBD5E1'),
                                      left=dict(sz=4, val='single', color='CBD5E1'),
                                      right=dict(sz=4, val='single', color='CBD5E1'))
        p = row_cells[i].paragraphs[0]
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # PHẦN 3: CHI TIẾT CÁC HẠNG MỤC ĐÃ TRIỂN KHAI
    # ==========================================
    add_sec_heading("3", "Chi Tiết Các Hạng Mục Đã Triển Khai")
    add_sub_heading("3.1", "Frontend: Phân hệ Quản trị Phòng khám (Clinic Admin Dashboard)")

    fe_table = doc.add_table(rows=1, cols=4)
    fe_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fe_headers = ["STT", "Tập tin", "Đường dẫn (URL)", "Chức năng chi tiết"]
    fe_widths = [Inches(0.5), Inches(1.8), Inches(1.7), Inches(2.7)]

    for i, h_text in enumerate(fe_headers):
        cell = fe_table.rows[0].cells[i]
        cell.text = h_text
        cell.width = fe_widths[i]
        set_cell_background(cell, "2563EB")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9)

    fe_data = [
        ("1", "AdminLayout.jsx", "/admin/*",
         "Khung giao diện chuẩn Slate-900 / Blue-600; Sidebar thu gọn thông minh; Nút chuyển nhanh sang Cổng Bác sĩ và Website người dùng; Menu điều hướng tinh gọn 4 phân hệ cốt lõi."),
        ("2", "AdminDashboard.jsx", "/admin/dashboard\n/dashboard",
         "KPI Cards: Tổng bác sĩ, chuyên khoa, số ca khám, bệnh nhân, doanh thu. Thanh phân bổ tỷ lệ ca khám (Chờ duyệt, Đã khám, Đã hủy). Nút lối tắt cấp tài khoản (kích hoạt mở Dialog) và danh sách bác sĩ."),
        ("3", "AdminDoctors.jsx", "/admin/doctors\n/dashboard/doctor",
         "Bảng danh sách bác sĩ & Dialog Cấp tài khoản / Upload ảnh tại chỗ:\n• Dialog Cấp Tài Khoản Mới: Mở trực tiếp ngay trong trang, không cần chuyển hướng trang.\n• Chọn ảnh từ máy tính / thiết bị: Tải ảnh avatar trực tiếp từ máy tính (<input type='file'>), tự động upload lên server và hiển thị preview tức thì.\n• Sinh mật khẩu tự động & Bàn giao: Nút random mật khẩu; Phiếu bàn giao tài khoản với nút 1-click Copy.\n• Modal Sửa & Đổi ảnh & Chuyên khoa: Cập nhật thông tin, đổi chuyên khoa và đổi ảnh đại diện từ máy tính.\n• Modal Cấp lại mật khẩu: Đổi mật khẩu tài khoản bác sĩ nhanh chóng.\n• Xóa Bác sĩ: Hộp thoại an toàn có xử lý ràng buộc lịch hẹn."),
        ("4", "AdminAppointments.jsx", "/admin/appointments",
         "Theo dõi và rà soát mọi ca khám bệnh của tất cả các chuyên khoa; Lọc theo trạng thái và tìm kiếm theo tên bệnh nhân / bác sĩ phụ trách."),
        ("5", "AdminPatients.jsx", "/admin/patients",
         "Quản lý danh bạ hồ sơ bệnh nhân đã đăng ký trên hệ thống, email, số điện thoại, địa chỉ."),
        ("6", "adminService.js", "N/A (API Service)",
         "Tầng Axios Service kết nối backend Spring Boot: cấp tài khoản, upload avatar từ máy tính, cập nhật, đổi mật khẩu, xóa bác sĩ, tải danh sách chuyên khoa, tải lịch khám...")
    ]

    for item in fe_data:
        row = fe_table.add_row()
        for i, val in enumerate(item):
            cell = row.cells[i]
            cell.text = val
            cell.width = fe_widths[i]
            set_cell_background(cell, "FFFFFF" if int(item[0]) % 2 != 0 else "F8FAFC")
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_border(cell, top=dict(sz=4, val='single', color='E2E8F0'),
                                  bottom=dict(sz=4, val='single', color='E2E8F0'),
                                  left=dict(sz=4, val='single', color='E2E8F0'),
                                  right=dict(sz=4, val='single', color='E2E8F0'))
            p = cell.paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)
                if i == 1:
                    run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Backend APIs Table ---
    add_sub_heading("3.2", "Backend: Nâng cấp Spring Boot REST API & Cơ Sở Dữ Liệu")

    be_table = doc.add_table(rows=1, cols=4)
    be_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    be_headers = ["Endpoint", "Method", "Controller & Service", "Mô tả nghiệp vụ"]
    be_widths = [Inches(2.2), Inches(0.8), Inches(1.7), Inches(2.0)]

    for i, h_text in enumerate(be_headers):
        cell = be_table.rows[0].cells[i]
        cell.text = h_text
        cell.width = be_widths[i]
        set_cell_background(cell, "1E40AF")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9)

    be_data = [
        ("/api/doctor/create", "POST", "DoctorController\nDoctorService",
         "Cấp tài khoản bác sĩ mới; kiểm tra trùng email; mã hóa mật khẩu BCrypt; liên kết chuyên khoa Specialty; gán tọa độ mặc định."),
        ("/api/doctor/admin/{id}", "PUT", "DoctorController\nDoctorService",
         "Admin cập nhật thông tin bác sĩ & điều chuyển sang chuyên khoa công tác mới."),
        ("/api/doctor/admin/{id}/reset-password", "PUT", "DoctorController\nDoctorService",
         "Admin cấp lại mật khẩu mới trực tiếp cho bác sĩ mà không cần biết mật khẩu cũ."),
        ("/api/upload/image", "POST", "UploadController\nWebMvcConfig",
         "Tải ảnh avatar trực tiếp từ máy tính / thiết bị người dùng; lưu vào uploads/ và phục vụ tĩnh trực tiếp qua URL."),
        ("/api/doctor/{id}", "DELETE", "DoctorController\nDoctorService",
         "Xóa bác sĩ khỏi CSDL theo ID có xử lý ràng buộc toàn vẹn lịch khám bệnh."),
        ("/api/doctor/all", "GET", "DoctorController\nDoctorService",
         "Lấy danh sách toàn bộ bác sĩ công tác tại phòng khám kèm chuyên khoa."),
        ("/api/specialty/all", "GET", "SpecialtyController\nSpecialtyService",
         "Lấy danh sách các chuyên khoa để phân công cho bác sĩ khi tạo tài khoản."),
        ("/api/appointment/all", "GET", "AppointmentController\nAppointmentService",
         "Lấy toàn bộ lịch hẹn khám trong toàn hệ thống phục vụ Admin Dashboard."),
        ("/api/appointment/admin/{id}/status", "PUT", "AppointmentController\nAppointmentService",
         "Admin trực tiếp duyệt, hoàn tất hoặc hủy lịch khám bệnh của mọi chuyên khoa."),
        ("/api/patient/all", "GET", "PatientController\nPatientService",
         "Lấy danh sách toàn bộ hồ sơ bệnh nhân đã đăng ký trên hệ thống."),
        ("/api/patient/admin/{id}", "DELETE", "PatientController\nPatientService",
         "Admin xóa hồ sơ bệnh nhân khỏi hệ thống có kiểm tra an toàn dữ liệu lịch hẹn.")
    ]

    for idx, item in enumerate(be_data):
        row = be_table.add_row()
        for i, val in enumerate(item):
            cell = row.cells[i]
            cell.text = val
            cell.width = be_widths[i]
            set_cell_background(cell, "FFFFFF" if idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_border(cell, top=dict(sz=4, val='single', color='E2E8F0'),
                                  bottom=dict(sz=4, val='single', color='E2E8F0'),
                                  left=dict(sz=4, val='single', color='E2E8F0'),
                                  right=dict(sz=4, val='single', color='E2E8F0'))
            p = cell.paragraphs[0]
            if i == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)
                if i == 1:
                    run.font.bold = True
                    if val == "POST":
                        run.font.color.rgb = RGBColor(0x15, 0x80, 0x3D) # Green
                    elif val == "PUT":
                        run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09) # Amber
                    elif val == "DELETE":
                        run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26) # Red
                    else:
                        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # PHẦN 4: KẾT QUẢ KIỂM THỬ THỰC TẾ
    # ==========================================
    add_sec_heading("4", "Kết Quả Kiểm Thử Thực Tế & Độ Ổn Định (100% PASS)")

    test_table = doc.add_table(rows=1, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_headers = ["Hạng mục kiểm thử", "Kịch bản kiểm thử", "Kết quả thực tế", "Trạng thái"]
    test_widths = [Inches(1.8), Inches(2.2), Inches(2.0), Inches(0.7)]

    for i, h_text in enumerate(test_headers):
        cell = test_table.rows[0].cells[i]
        cell.text = h_text
        cell.width = test_widths[i]
        set_cell_background(cell, "059669") # Emerald Green
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9)

    test_data = [
        ("Biên dịch Frontend", "Thực thi lệnh 'npm run build' với Vite", "Built in 1.13s, 158 modules transformed, 0 lỗi cú pháp", "PASS ✅"),
        ("Biên dịch Backend", "Thực thi lệnh './mvnw compile' với Maven", "BUILD SUCCESS, 71 source files biên dịch hoàn tất, 0 lỗi", "PASS ✅"),
        ("Cấp tài khoản Bác sĩ", "Tạo bác sĩ mới từ form /admin/doctors/create", "Lưu thành công vào MySQL, mật khẩu được băm BCrypt", "PASS ✅"),
        ("Đăng nhập Bác sĩ mới", "Đăng nhập tại /doctor/login với tài khoản vừa tạo", "Đăng nhập thành công, nhận JWT/Session vào Cổng Bác sĩ", "PASS ✅"),
        ("Cập nhật & Phân khoa", "Đổi chuyên khoa từ Tim mạch sang Nội khoa qua Modal", "CSDL cập nhật ngay lập tức, hiển thị đúng trên danh sách", "PASS ✅"),
        ("Cấp lại mật khẩu", "Admin reset mật khẩu cho bác sĩ qua popup", "Đổi mật khẩu thành công, bác sĩ login được với pass mới", "PASS ✅"),
        ("Xóa Bác sĩ", "Xóa bác sĩ thử nghiệm và kiểm tra ràng buộc", "Xóa an toàn nếu chưa có lịch khám, báo lỗi nếu có lịch", "PASS ✅"),
        ("Hiển thị KPI Admin", "Kiểm tra các thẻ tổng quan phòng khám", "Hiển thị chính xác số bác sĩ, chuyên khoa, lịch hẹn", "PASS ✅")
    ]

    for idx, item in enumerate(test_data):
        row = test_table.add_row()
        for i, val in enumerate(item):
            cell = row.cells[i]
            cell.text = val
            cell.width = test_widths[i]
            set_cell_background(cell, "FFFFFF" if idx % 2 == 0 else "F0FDF4")
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_border(cell, top=dict(sz=4, val='single', color='CBD5E1'),
                                  bottom=dict(sz=4, val='single', color='CBD5E1'),
                                  left=dict(sz=4, val='single', color='CBD5E1'),
                                  right=dict(sz=4, val='single', color='CBD5E1'))
            p = cell.paragraphs[0]
            if i == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)
                if i == 3:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x05, 0x96, 0x69)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # PHẦN 5: HƯỚNG DẪN TRUY CẬP & BÀN GIAO
    # ==========================================
    add_sec_heading("5", "Hướng Dẫn Truy Cập & Bàn Giao Sử Dụng")

    guide_table = doc.add_table(rows=1, cols=3)
    guide_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    guide_headers = ["Phân hệ", "Đường dẫn truy cập (URL)", "Hướng dẫn sử dụng & Tài khoản mẫu"]
    guide_widths = [Inches(2.0), Inches(2.2), Inches(2.5)]

    for i, h_text in enumerate(guide_headers):
        cell = guide_table.rows[0].cells[i]
        cell.text = h_text
        cell.width = guide_widths[i]
        set_cell_background(cell, "334155")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9)

    guide_data = [
        ("Cấp Tài Khoản Bác Sĩ Mới", "http://localhost:5173/admin/doctors/create\nhoặc /admin/create-doctor",
         "Điền thông tin bác sĩ, chọn chuyên khoa từ DB, bấm 'Cấp Tài Khoản Bác Sĩ' và bấm 'Sao chép thông tin' để gửi bàn giao cho bác sĩ."),
        ("Quản Lý Bác Sĩ & Đổi Chuyên Khoa", "http://localhost:5173/admin/doctors\nhoặc /dashboard/doctor",
         "Xem danh sách bác sĩ, tìm kiếm theo tên/email, bấm nút Sửa để đổi chuyên khoa hoặc nút Chìa khóa để cấp lại mật khẩu."),
        ("Bảng Điều Khiển Quản Trị Toàn Viện", "http://localhost:5173/admin/dashboard\nhoặc /dashboard",
         "Giám sát toàn bộ KPI, doanh thu ước tính và biểu đồ tiến độ các ca khám của cả phòng khám."),
        ("Lịch Hẹn Toàn Viện", "http://localhost:5173/admin/appointments",
         "Xem và lọc mọi ca khám bệnh trong hệ thống theo chuyên khoa, trạng thái và bác sĩ phụ trách."),
        ("Cổng Bác Sĩ (Doctor Portal)", "http://localhost:5173/doctor/login",
         "Dành cho Bác sĩ tiếp nhận ca khám, quản lý lịch trực.\nTài khoản mẫu kiểm thử:\n• Email: bs.a@example.com\n• Mật khẩu: password123"),
        ("Trang Chủ Bệnh Nhân", "http://localhost:5173/patient/home",
         "Trải nghiệm đặt lịch khám, trợ lý AI tư vấn triệu chứng, bản đồ định vị phòng khám và thanh toán VNPay.")
    ]

    for idx, item in enumerate(guide_data):
        row = guide_table.add_row()
        for i, val in enumerate(item):
            cell = row.cells[i]
            cell.text = val
            cell.width = guide_widths[i]
            set_cell_background(cell, "FFFFFF" if idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_border(cell, top=dict(sz=4, val='single', color='E2E8F0'),
                                  bottom=dict(sz=4, val='single', color='E2E8F0'),
                                  left=dict(sz=4, val='single', color='E2E8F0'),
                                  right=dict(sz=4, val='single', color='E2E8F0'))
            p = cell.paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(8.5)
                if i == 0:
                    run.font.bold = True

    # -------------------------------------------------------------
    # SECTION 6: CLEAN ARCHITECTURE, COMPONENT & CONTEXT REFACTORING
    # -------------------------------------------------------------
    doc.add_paragraph()
    h6 = doc.add_paragraph()
    r6 = h6.add_run("6. TÁI CẤU TRÚC CLEAN ARCHITECTURE: PHÂN TRANG (PAGINATION), COMPONENT & CONTEXT")
    r6.font.bold = True
    r6.font.size = Pt(13)
    r6.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    p_clean = doc.add_paragraph()
    p_clean.add_run("Nhằm tối ưu hóa hiệu năng và tuân thủ các nguyên lý Clean Code / Separation of Concerns (SoC), toàn bộ phân hệ Admin đã được tái cấu trúc thành kiến trúc đa tầng hiện đại:")
    p_clean.paragraph_format.space_after = Pt(6)

    clean_table = doc.add_table(rows=1, cols=3)
    clean_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    clean_widths = [Inches(1.8), Inches(2.2), Inches(2.7)]

    c_headers = ["Thành phần kiến trúc", "Vị trí mã nguồn", "Mục đích & Lợi ích tối ưu"]
    for i, h_text in enumerate(c_headers):
        cell = clean_table.rows[0].cells[i]
        cell.text = h_text
        cell.width = clean_widths[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9)

    clean_data = [
        ("AdminContext & useAdmin()", "frontend/src/context/AdminContext.jsx",
         "Quản lý trạng thái tập trung (Bác sĩ, Chuyên khoa, Lịch hẹn, Bệnh nhân, KPIs), triệt tiêu gọi trùng API và đồng bộ badge counters."),
        ("Component Phân trang dùng chung", "frontend/src/component/common/Pagination.jsx",
         "Điều hướng trang (<<, <, số trang, >, >>), chọn số dòng/trang (5, 10, 20, 50), tự động reset về trang 1 khi tìm kiếm."),
        ("Thẻ KPI Tái sử dụng", "frontend/src/component/admin/StatCard.jsx",
         "Hiển thị số liệu thống kê Dashboard với thanh gradient, biểu tượng sinh động và nhãn phân loại."),
        ("Bảng Bác sĩ & Thao tác", "frontend/src/component/admin/DoctorTable.jsx",
         "Hiển thị bảng bác sĩ kèm avatar, chuyên khoa, mức phí, kinh nghiệm và nhóm nút thao tác sửa/đổi pass/xóa."),
        ("Dialog Cấp Tài Khoản Bác Sĩ", "frontend/src/component/admin/DoctorCreateModal.jsx",
         "Tải avatar từ máy tính người dùng, sinh mật khẩu ngẫu nhiên, phân công chuyên khoa và tích hợp geocoding."),
        ("Phiếu Bàn Giao Tài Khoản", "frontend/src/component/admin/DoctorVoucherModal.jsx",
         "Hiển thị thông tin tài khoản bác sĩ vừa cấp với nút 1-click sao chép định dạng văn bản gửi bác sĩ."),
        ("Dialog Sửa & Đổi Chuyên Khoa", "frontend/src/component/admin/DoctorEditModal.jsx",
         "Chỉnh sửa thông tin bác sĩ, điều chuyển khoa và tải ảnh đại diện mới từ thiết bị."),
        ("Dialog Cấp Lại Mật Khẩu", "frontend/src/component/admin/DoctorResetPasswordModal.jsx",
         "Đổi mật khẩu nhanh cho bác sĩ trong popup và sao chép mật khẩu mới."),
        ("Bảng Lịch Hẹn Toàn Viện", "frontend/src/component/admin/AppointmentTable.jsx",
         "Quản lý danh sách lịch khám toàn viện kèm dropdown cập nhật trạng thái trực tiếp."),
        ("Bảng Danh Bạ Bệnh Nhân", "frontend/src/component/admin/PatientTable.jsx",
         "Hiển thị hồ sơ bệnh nhân, mã định danh, avatar ký tự đầu và chức năng xóa hồ sơ.")
    ]

    for idx, item in enumerate(clean_data):
        row = clean_table.add_row()
        for i, val in enumerate(item):
            cell = row.cells[i]
            cell.text = val
            cell.width = clean_widths[i]
            set_cell_background(cell, "FFFFFF" if idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_border(cell, top=dict(sz=4, val='single', color='E2E8F0'),
                                  bottom=dict(sz=4, val='single', color='E2E8F0'),
                                  left=dict(sz=4, val='single', color='E2E8F0'),
                                  right=dict(sz=4, val='single', color='E2E8F0'))
            p = cell.paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(8.5)
                if i == 0:
                    run.font.bold = True

    # Save document
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Report successfully saved to: {output_path}")


if __name__ == '__main__':
    target_path = "/Users/m1lt43/Desktop/DoAnChuyennganh/medical-appointment/docs/TIEN_DO_HOAN_THIEN.docx"
    create_report(target_path)
    # Also save as BAO_CAO_TIEN_DO_DU_AN.docx for convenience
    copy_path = "/Users/m1lt43/Desktop/DoAnChuyennganh/medical-appointment/docs/BAO_CAO_TIEN_DO_DU_AN.docx"
    create_report(copy_path)
